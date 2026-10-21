#!/usr/bin/env python3
"""Import the verified Kanei customer workbook into the scoped customer master."""
import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.request
import re
from pathlib import Path

try:
    import openpyxl
except ImportError as error:
    raise SystemExit("openpyxl is required: pip install openpyxl") from error


def request_json(url: str, method: str, payload: dict | None = None, cookie: str | None = None):
    body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if cookie:
        headers["Cookie"] = cookie
    request = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            content = response.read().decode("utf-8")
            return response.status, json.loads(content) if content else {}, response.headers
    except urllib.error.HTTPError as error:
        content = error.read().decode("utf-8", errors="replace")
        try:
            parsed = json.loads(content)
        except json.JSONDecodeError:
            parsed = {"raw": content}
        return error.code, parsed, error.headers


def clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def login(base_url: str, email: str, password: str) -> str:
    status, payload, headers = request_json(
        f"{base_url}/api/auth/login", "POST", {"email": email, "password": password}
    )
    if status != 200:
        raise SystemExit(f"login failed ({status}): {payload}")
    set_cookie = headers.get("Set-Cookie", "")
    return set_cookie.split(";", 1)[0]


def rows_from_workbook(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "顧客マスター" not in workbook.sheetnames:
        raise SystemExit("workbook does not contain 顧客マスター")
    worksheet = workbook["顧客マスター"]
    for row in worksheet.iter_rows(min_row=5, values_only=True):
        number = clean(row[0])
        if not number or not number.isdigit():
            continue
        number_value = int(number)
        public_name = clean(row[2]) if len(row) > 2 else ""
        original_name = clean(row[1]) if len(row) > 1 else ""
        postal = clean(row[3]) if len(row) > 3 else ""
        address = clean(row[4]) if len(row) > 4 else ""
        public_tel = clean(row[5]) if len(row) > 5 else ""
        old_tel = clean(row[6]) if len(row) > 6 else ""
        contact = clean(row[9]) if len(row) > 9 else ""
        contact_email = clean(row[10]) if len(row) > 10 else ""
        public_email = clean(row[11]) if len(row) > 11 else ""
        name = public_name or original_name
        if not name:
            continue
        public_email = public_email if EMAIL_RE.fullmatch(public_email) else ""
        contact_email = contact_email if EMAIL_RE.fullmatch(contact_email) else ""
        yield {
            "customerCode": f"K-{number_value:04d}",
            "customerName": name,
            "customerPostalCode": postal,
            "customerAddress": address,
            "customerContact": contact,
            "customerTelephone": public_tel or old_tel,
            "customerEmail": public_email or contact_email,
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--base-url", default=os.environ.get("BASE_URL", "http://127.0.0.1:3002"))
    parser.add_argument("--email", default=os.environ.get("IMPORT_EMAIL", os.environ.get("ADMIN_EMAIL", "")))
    parser.add_argument("--password", default=os.environ.get("IMPORT_PASSWORD", os.environ.get("ADMIN_PASSWORD", "")))
    parser.add_argument("--cookie", default=os.environ.get("IMPORT_COOKIE", ""))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not args.cookie and (not args.email or not args.password):
        raise SystemExit("IMPORT_EMAIL and IMPORT_PASSWORD are required")

    cookie = args.cookie or login(args.base_url, args.email, args.password)
    imported = skipped = errors = 0
    for customer in rows_from_workbook(Path(args.workbook)):
        if args.dry_run:
            print(f"DRY {customer['customerCode']} {customer['customerName']}")
            imported += 1
            continue
        status, payload, _ = request_json(
            f"{args.base_url}/api/customers", "POST", customer, cookie
        )
        if status in (200, 201):
            imported += 1
            print(f"OK {customer['customerCode']} {customer['customerName']}")
        elif status == 500 and payload.get("error") == "customer_save_failed":
            skipped += 1
            print(f"SKIP duplicate {customer['customerCode']} {customer['customerName']}")
        else:
            errors += 1
            print(f"ERROR {customer['customerCode']} ({status}): {payload}", file=sys.stderr)
        time.sleep(0.03)

    print(f"done: imported={imported} skipped={skipped} errors={errors}")
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
